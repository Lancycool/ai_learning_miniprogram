from pathlib import Path

import pytest
from pypdf import PdfWriter
from pypdf.generic import ArrayObject, DictionaryObject, NameObject, TextStringObject, DecodedStreamObject
from docx import Document

from app.core.config import Settings
from app.core.knowledge_errors import KnowledgeError


def pdf_sample(path, *, encrypted=False, blank=False, image=False):
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    if not blank:
        descendant = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/CIDFontType0"), NameObject("/BaseFont"): NameObject("/STSong-Light"),
            NameObject("/CIDSystemInfo"): DictionaryObject({NameObject("/Registry"): TextStringObject("Adobe"),
                NameObject("/Ordering"): TextStringObject("GB1"), NameObject("/Supplement"): __import__('pypdf').generic.NumberObject(0)})})
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type0"),
            NameObject("/BaseFont"): NameObject("/STSong-Light"), NameObject("/Encoding"): NameObject("/UniGB-UCS2-H"),
            NameObject("/DescendantFonts"): ArrayObject([writer._add_object(descendant)])})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        rows = ["第一章 客服流程", "客服先核对订单。", "第二章 退款流程", "退款由主管确认。"]
        stream.set_data(("BT /F1 12 Tf 40 750 Td " + " 0 -24 Td ".join(f"<{s.encode('utf-16be').hex()}> Tj" for s in rows) + " ET").encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
        writer.add_outline_item("第一章 客服流程", 0)
        if image:
            image_object = DecodedStreamObject()
            image_object.set_data(b'\xff\x00\x00')
            image_object.update({NameObject('/Type'): NameObject('/XObject'), NameObject('/Subtype'): NameObject('/Image')})
            page['/Resources'][NameObject('/XObject')] = DictionaryObject({NameObject('/Image1'): writer._add_object(image_object)})
    if encrypted:
        writer.encrypt("synthetic-password")
    writer.write(path)


@pytest.mark.parametrize("extension", ["pdf", "docx", "md", "txt"])
async def test_real_loaders_preserve_chinese_and_chapter_locations(tmp_path, extension):
    from app.services.knowledge_parser import KnowledgeParser
    path = tmp_path / f"training.{extension}"
    if extension == "pdf":
        pdf_sample(path)
    elif extension == "docx":
        doc = Document()
        doc.add_paragraph("培训介绍")
        doc.add_heading("第一章 客服流程", level=1)
        doc.add_paragraph("客服先核对订单。")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "客服", "核对订单"
        doc.add_heading("第二章 退款流程", level=1)
        doc.add_paragraph("退款由主管确认。")
        doc.save(path)
    else:
        prefix = "# " if extension == "md" else ""
        path.write_text(f"培训介绍\n{prefix}第一章 客服流程\n客服先核对订单。\n{prefix}第二章 退款流程\n退款由主管确认。", encoding="utf-8")
    parsed = await KnowledgeParser(Settings(_env_file=None)).parse(path, extension)
    assert "客服先核对订单。" in parsed["text"] and "退款由主管确认。" in parsed["text"]
    assert parsed["loader"] == {"pdf": "PyPDFLoader", "docx": "Docx2txtLoader", "md": "UnstructuredMarkdownLoader", "txt": "TextLoader"}[extension]
    assert len(parsed["chapters"]) >= 2
    for chapter in parsed["chapters"]:
        assert 0 <= chapter["start_offset"] < chapter["end_offset"] <= len(parsed["text"])
    assert any("第一章" in c["title"] for c in parsed["chapters"])
    if extension in ("docx", "md", "txt"):
        assert parsed["chapters"][0]["title"] == "未分章内容"
    if extension == "docx":
        assert "核对订单" in parsed["text"]


async def test_nested_markdown_and_unassigned_content_keep_full_text(tmp_path):
    from app.services.knowledge_parser import KnowledgeParser
    path = tmp_path / "chapters.md"
    original = "开头\n# 主章\n主章介绍\n## 子章\n子章正文\n# 末章\n结束\n"
    path.write_text(original, encoding="utf-8")
    parsed = await KnowledgeParser(Settings(_env_file=None)).parse(path, "md")
    assert parsed["text"] == original
    assert [c["level"] for c in parsed["chapters"]] == [1, 1, 2, 1]
    child = parsed["chapters"][2]
    assert parsed["chapters"][1]["end_offset"] >= child["end_offset"]


async def test_text_pdf_with_images_discloses_unparsed_image_content(tmp_path):
    from app.services.knowledge_parser import KnowledgeParser
    path = tmp_path / 'image-training.pdf'
    pdf_sample(path, image=True)
    parsed = await KnowledgeParser(Settings(_env_file=None)).parse(path, 'pdf')
    assert '核对订单' in parsed['text']
    assert any(w['code'] == 'embedded_images' for w in parsed['warnings'])


@pytest.mark.parametrize("kind,reason", [("encrypted", "encrypted_pdf"), ("blank", "scanned_pdf"), ("broken", "invalid_file")])
async def test_unsupported_pdf_is_controlled(tmp_path, kind, reason):
    from app.services.knowledge_parser import KnowledgeParser
    path = tmp_path / "problem.pdf"
    if kind == "broken":
        path.write_bytes(b"%PDF-1.7 broken synthetic document")
    else:
        pdf_sample(path, encrypted=kind == "encrypted", blank=kind == "blank")
    with pytest.raises(KnowledgeError) as caught:
        await KnowledgeParser(Settings(_env_file=None)).parse(path, "pdf")
    assert caught.value.reason == reason
    assert str(path) not in caught.value.public_message


async def test_parse_limits_and_real_child_timeout(tmp_path):
    from app.services.knowledge_parser import KnowledgeParser
    path = tmp_path / "limits.txt"
    path.write_text("资料内容" * 20, encoding="utf-8")
    settings = Settings(_env_file=None, KNOWLEDGE_MAX_CHARACTERS=10)
    with pytest.raises(KnowledgeError) as caught:
        await KnowledgeParser(settings).parse(path, "txt")
    assert caught.value.reason == "text_limit"
    settings.knowledge_parse_timeout_seconds = 0.001
    with pytest.raises(KnowledgeError) as caught:
        await KnowledgeParser(settings).parse(path, "txt")
    assert caught.value.reason == "parse_timeout"


async def test_parser_cancel_terminates_child(tmp_path):
    import asyncio
    from app.services.knowledge_parser import KnowledgeParser
    path = tmp_path / "cancel.md"
    path.write_text("# 资料\n正文", encoding="utf-8")
    parser = KnowledgeParser(Settings(_env_file=None))
    task = asyncio.create_task(parser.parse(path, "md"))
    await asyncio.sleep(0.05)
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert not parser.active_processes


async def test_images_are_disclosed_and_pdf_page_limit_is_enforced(tmp_path):
    from app.services.knowledge_parser import KnowledgeParser
    path = tmp_path / "images.md"
    path.write_text("# 原题\n1. 请看下图。\n![图](https://example.invalid/private.png)\n", encoding="utf-8")
    parsed = await KnowledgeParser(Settings(_env_file=None)).parse(path, "md")
    assert parsed["warnings"][0]["code"] == "embedded_images"
    assert "https://example.invalid/private.png" in parsed["text"]
    writer = PdfWriter()
    writer.add_blank_page(width=600, height=800)
    writer.add_blank_page(width=600, height=800)
    path = tmp_path / "too-many.pdf"
    writer.write(path)
    with pytest.raises(KnowledgeError) as caught:
        await KnowledgeParser(Settings(_env_file=None, KNOWLEDGE_MAX_PAGES=1)).parse(path, "pdf")
    assert caught.value.reason == "page_limit"
