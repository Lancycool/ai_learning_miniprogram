"""Cross-layer checks use real loaders, MySQL and persistent Chroma.

Only the external embedding HTTP boundary is replaced with a fixed 1024-vector.
"""
import json
import httpx
import pytest
from pydantic import SecretStr
from docx import Document
from app.models.knowledge import KnowledgeScope
from app.services.bailian_embeddings import BailianEmbeddings
from app.services.knowledge_processing_service import KnowledgeProcessor
from app.services.knowledge_service import KnowledgeService
from app.services.knowledge_task_service import KnowledgeTaskWorker
from app.services.knowledge_vector_store import KnowledgeVectorStore
from tests.test_knowledge_api import knowledge_env, client, create_base, upload
from tests.test_knowledge_parser import pdf_sample


async def test_four_uploads_chapters_persisted_index_rebuild_and_cleanup(knowledge_env, tmp_path):
    env = knowledge_env
    env[2].bailian_api_key = SecretStr('synthetic-only')
    def embedding_response(request):
        payload = json.loads(request.content)
        return httpx.Response(200, json={'output': {'embeddings': [
            {'text_index': i, 'embedding': [1.0]+[0.0]*1023}
            for i, _ in enumerate(payload['input']['texts'])]}})
    async with httpx.AsyncClient(transport=httpx.MockTransport(embedding_response)) as upstream:
        embedding = BailianEmbeddings(env[2], async_client=upstream)
        store = KnowledgeVectorStore(env[2], embedding)
        runner = KnowledgeTaskWorker(env[1], env[2], processor=KnowledgeProcessor(env[1], env[2], vector_store=store))
        try:
            async with client() as api:
                base_id = await create_base(api)
                pdf = tmp_path/'training.pdf'; pdf_sample(pdf)
                word = tmp_path/'training.docx'; document = Document()
                document.add_heading('第一章 客服流程', level=1); document.add_paragraph('客服先核对订单。')
                document.add_heading('第二章 退款流程', level=1); document.add_paragraph('退款由主管确认。'); document.save(word)
                text = '第一章 客服流程\n客服先核对订单。\n第二章 退款流程\n退款由主管确认。'
                files = [('pdf', pdf.read_bytes()), ('docx', word.read_bytes()),
                    ('md', text.replace('第一章', '# 第一章').replace('第二章', '# 第二章').encode()), ('txt', text.encode())]
                ids = []
                for extension, content in files:
                    task = (await upload(api, base_id, content, f'training.{extension}', f'acceptance-upload-{extension}-0001')).json()['data']
                    assert task['status'] == 'queued' and task['result'] is None
                    await runner.run_once()
                    done = (await api.get(f"/api/v1/knowledge-processing-tasks/{task['task_id']}")).json()['data']
                    assert done['status'] == 'succeeded' and done['document']['index_status'] == 'ready'
                    doc_id = done['document']['document_id']; ids.append(doc_id)
                    preview = (await api.get(f'/api/v1/knowledge-documents/{doc_id}/preview')).json()['data']
                    assert len(preview['chapters']) == 2 and '客服先核对订单。' in preview['text']
                    selected = KnowledgeScope(knowledge_base_id=base_id, documents=[{'document_id': doc_id,
                        'version_id': done['document']['version_id'], 'chapter_ids': [preview['chapters'][0]['chapter_id']]}])
                    async with env[1]() as db:
                        scope = await KnowledgeService(db, env[2]).resolve_scope(env[3], selected)
                    reopened = KnowledgeVectorStore(env[2], embedding)
                    hits = await reopened.retrieve('客服核对订单', env[3].id, base_id, scope)
                    assert hits and all('退款' not in hit['text'] for hit in hits)
                # A corrected chapter creates a new version and a queued index.
                changed = await api.patch(f'/api/v1/knowledge-documents/{ids[0]}/chapters', json={
                    'version_id': (await api.get(f'/api/v1/knowledge-documents/{ids[0]}')).json()['data']['version_id'],
                    'chapters': [{'title': '合并章节', 'start_offset': 0, 'end_offset': len(text)}]})
                assert changed.status_code == 202
                await runner.run_once()
                assert (await api.get(f'/api/v1/knowledge-documents/{ids[0]}/preview')).json()['data']['chapters'][0]['title'] == '合并章节'
                deleted = (await api.delete(f'/api/v1/knowledge-bases/{base_id}')).json()['data']
                assert len(deleted['cleanup_task_ids']) == 4
                assert (await api.get('/api/v1/knowledge-processing-tasks/cleanup')).json()['data']['total'] == 4
                for _ in ids: await runner.run_once()
                assert (await api.get('/api/v1/knowledge-processing-tasks/cleanup')).json()['data']['total'] == 0
                root = env[2].private_directory(env[2].knowledge_storage_directory)
                assert not [p for p in root.rglob('*') if p.is_file()]
                chroma = next(iter(store.stores.values()))._client
                assert all(chroma.get_collection(collection.name).count() == 0 for collection in chroma.list_collections())
        finally:
            if store.stores: next(iter(store.stores.values()))._client._system.stop()
